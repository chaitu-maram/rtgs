from django.shortcuts import render
from django.http import HttpResponse, JsonResponse
from .forms import AxisBankForm, UnionBankForm
from .models import AxisBankRemitter, UnionBankRemitter, AxisBankBeneficiary
from io import BytesIO
from docxtpl import DocxTemplate
from django.conf import settings
from num2words import num2words
import os
from datetime import datetime
from utils.gsheets import append_row_to_sheet
from django.contrib.auth import authenticate, login, logout


import os
from io import BytesIO
from django.http import HttpResponse, JsonResponse
from django.shortcuts import render
from django.conf import settings
from docxtpl import DocxTemplate
from num2words import num2words
from .models import AxisBankRemitter, AxisBankBeneficiary
from .forms import AxisBankForm






from django.shortcuts import render
from django.http import HttpResponse, JsonResponse
from .forms import AxisBankForm
from .models import AxisBankRemitter, AxisBankBeneficiary
from .models import RemitterAccount
from docxtpl import DocxTemplate
from io import BytesIO
from django.conf import settings
from num2words import num2words
import os


def axis_bank(request):
    # ✅ Fetch remitter details using unique_code from session
    unique_code = request.session.get('unique_code', None)
    remitter_data = {}

    if unique_code:
        remitter_acc = RemitterAccount.objects.filter(unique_code=unique_code, bank_name__icontains='Axis').first()
        if remitter_acc:
            remitter_data = {
                'remitter_name': remitter_acc.remitter_name or "",
                'remitter_account_number': remitter_acc.remitter_account_number or "",
                'remitter_address': remitter_acc.remitter_address or "",
                'branch': remitter_acc.branch or "",
                'pan_no': remitter_acc.pan_no or "",
                'mobile_number': remitter_acc.mobile_number or "",
                'other_number': remitter_acc.other_number or "",
                'remitter_ifsc_code': remitter_acc.ifsc_code or "",
            }

    if request.method == 'POST':
        form = AxisBankForm(request.POST)
        if form.is_valid():
            data = form.cleaned_data



            # --- Save remitter if not already saved in RemitterAccount ---
            remitter_account = data.get("remitter_account_number", "")
            bank_name = data.get("remitter_branch") or data.get("bank_name") or "Axis Bank"  # fallback

            existing_remitter = RemitterAccount.objects.filter(remitter_account_number=remitter_account, bank_name__icontains=bank_name).first()

            if not existing_remitter:
                RemitterAccount.objects.create(
                    unique_code=request.session.get('unique_code', ''),
                    bank_name=bank_name,
                    remitter_account_number=remitter_account,
                    remitter_name=data.get("remitter_name", ""),
                    ifsc_code=data.get("remitter_ifsc_code", ""),
                    remitter_address=data.get("remitter_address", ""),
                    branch=data.get("remitter_branch", ""),
                    pan_no=data.get("pan_no", ""),
                    mobile_number=data.get("mobile_number", ""),
                    other_number=data.get("other_number", ""),
                )

            # --- Save beneficiary if new ---
            beneficiary_account = data.get("beneficiary_account_number", "")
            existing_beneficiary = AxisBankBeneficiary.objects.filter(account_number=beneficiary_account).first()

            if not existing_beneficiary:
                AxisBankBeneficiary.objects.create(
                    unique_code=request.session.get("unique_code", ""),
                    # unique_code=unique_code,
                    name=data.get("beneficiary_name", ""),
                    account_number=beneficiary_account,
                    bank_name=data.get("beneficiary_bank_name", ""),
                    bank_address=data.get("beneficiary_bank_address", ""),
                    # lei_code=data.get("beneficiary_LEI_code", ""),
                    ifsc_code=data.get("beneficiary_bank_ifsc_code", ""),
                    # transaction_amount=data.get("transaction_amount", 0),
                )

            # --- Prepare Word document ---
            template_path = os.path.join(settings.MEDIA_ROOT, 'Axis Bank Rtgs Form.docx')
            doc = DocxTemplate(template_path)

            amount_value = data.get("transaction_amount", 0)
            amount_words = num2words(amount_value, to="currency", lang="en_IN", currency="INR").upper()

            # --- Safe rendering for all fields ---
            context = {
                "date": data.get("date").strftime("%d-%m-%Y") if data.get("date") else "",
                "branch": data.get("branch", ""),
                "amount": str(amount_value),
                "amount_words": amount_words,
                "pan_number_list": list(data.get("pan_no") or ""),
                "remitter_account_number": remitter_account,
                "remitter_account_number_list": list(remitter_account or ""),
                "remitter_name_list": list(data.get("remitter_name") or ""),
                "remitter_address_list": list(data.get("remitter_address") or ""),
                "remitter_LEI_code": data.get("remitter_LEI_code") or "",
                "remitter_LEI_code_list": list(data.get("remitter_LEI_code") or ""),
                "beneficiary_LEI_code": data.get("beneficiary_LEI_code") or "",
                "l": list(data.get("beneficiary_LEI_code") or ""),
                "beneficiary_bank_address": data.get("beneficiary_bank_address") or "",
                "a": list(data.get("beneficiary_bank_address") or ""),
                "cheque_number_list": list(data.get("cheque_number") or ""),
                "cheque_date_list": list(data.get("cheque_date").strftime("%d%m%Y") if data.get("cheque_date") else ""),
                "mobile_number": data.get("mobile_number") or "",
                "mobile_number_list": list(data.get("mobile_number") or ""),
                "other_number": data.get("other_number") or "",
                "other_number_list": list(data.get("other_number") or ""),
                "beneficiary_name": data.get("beneficiary_name") or "",
                "beneficiary_name_list": list(data.get("beneficiary_name") or ""),
                "beneficiary_account_number": beneficiary_account,
                "b": list(beneficiary_account or ""),
                "beneficiary_bank_name": data.get("beneficiary_bank_name") or "",
                "n": list(data.get("beneficiary_bank_name") or ""),
                "beneficiary_bank_ifsc_code": data.get("beneficiary_bank_ifsc_code") or "",
                "sc": list(data.get("beneficiary_bank_ifsc_code") or ""),
                "sender_to_receiver_information": data.get("sender_to_receiver_information") or "",
                "sr": list(data.get("sender_to_receiver_information") or ""),
            }

            doc.render(context)
            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)

            response = HttpResponse(
                buffer.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )
            response['Content-Disposition'] = 'attachment; filename="AxisBankForm_Filled.docx"'
            response.set_cookie('fileDownload', 'true', max_age=10, path='/')
            return response

    else:
        form = AxisBankForm(initial=remitter_data)

    return render(request, 'form_filler/axis_bank.html', {'form': form})



from django.http import JsonResponse
from .models import AxisBankBeneficiary

def fetch_beneficiary(request):
    q = request.GET.get('q', '').strip()
    unique_code = request.session.get('unique_code', '').strip()

    print(f"\n🧠 SESSION unique_code = '{unique_code}'")  # Debug 1
    print(f"🔎 Searching for query = '{q}'")

    if not q or not unique_code:
        print("⚠️ Missing query or unique_code — returning empty results.")
        return JsonResponse({'results': []})

    # ✅ Build the queryset step-by-step
    queryset = AxisBankBeneficiary.objects.filter(name__icontains=q)
    print(f"📦 Initial matches for name: {queryset.count()}")

    queryset = queryset.filter(unique_code=unique_code)
    print(f"🧮 After filtering by unique_code={unique_code}: {queryset.count()} matches")

    queryset = queryset.exclude(unique_code__isnull=True).exclude(unique_code="")[:10]
    print(f"✅ Final results count (after exclude): {queryset.count()}")

    results = [
        {
            'name': b.name,
            'account_number': b.account_number,
            'bank_name': b.bank_name,
            'bank_address': b.bank_address,
            'ifsc_code': b.ifsc_code,
        }
        for b in queryset
    ]

    print(f"📤 Returned {len(results)} results for {unique_code}\n")
    return JsonResponse({'results': results})






# ✅ Fetch remitter list
def fetch_remitter(request):
    q = request.GET.get('q', '').strip()
    if not q:
        return JsonResponse({"results": []})

    remitters = AxisBankRemitter.objects.filter(account_number__icontains=q)[:10]
    results = [
        {
            "account_number": r.account_number,
            "name": r.name,
            "ifsc_code": r.ifsc_code,
            "address": r.address,
            "cheque_number": r.cheque_number,
            "mobile_number": r.mobile_number,
            "other_number": r.other_number,
            "branch": r.branch,
            "pan_no": r.pan_no,
        }
        for r in remitters
    ]
    return JsonResponse({"results": results})








# from django.shortcuts import render
# from django.http import HttpResponse
# from io import BytesIO
# import os
# from num2words import num2words
# from docxtpl import DocxTemplate
# from .forms import UnionBankForm
# from .models import RemitterAccount, AxisBankBeneficiary
# from django.conf import settings

# def union_bank(request):
#     unique_code = request.session.get('unique_code')
#     remitter_data = {}

#     # Prefill form if remitter exists
#     if unique_code:
#         remitter_acc = RemitterAccount.objects.filter(
#             unique_code=unique_code,
#             bank_name__icontains="Union Bank"
#         ).first()

#         if remitter_acc:
#             remitter_data = {
#                 "remitter_name": remitter_acc.remitter_name or "",
#                 "remitter_account_number": remitter_acc.remitter_account_number or "",
#                 "remitter_address": remitter_acc.remitter_address or "",
#                 "branch": remitter_acc.branch or "",
#                 "pan_number": remitter_acc.pan_no or "",
#                 "mobile_number": remitter_acc.mobile_number or "",
#                 "other_number": remitter_acc.other_number or "",
#             }

#     if request.method == "POST":
#         form = UnionBankForm(request.POST)
#         if form.is_valid():
#             data = form.cleaned_data

#             # Save remitter if new
#             existing_remitter = RemitterAccount.objects.filter(
#                 remitter_account_number=data.get("remitter_account_number"),
#                 bank_name__icontains="Union Bank"
#             ).first()

#             if not existing_remitter:
#                 RemitterAccount.objects.create(
#                     unique_code=unique_code,
#                     bank_name="Union Bank",
#                     remitter_account_number=data.get("remitter_account_number"),
#                     remitter_name=data.get("remitter_name", ""),
#                     remitter_address=data.get("remitter_address", ""),
#                     mobile_number=data.get("mobile_number", ""),
#                     pan_no=data.get("pan_number", ""),
#                     branch=data.get("branch", ""),
#                     other_number=data.get("other_number", "")
#                 )


#             # Save beneficiary if new
#             beneficiary_account = data.get("beneficiary_account_number")
#             existing_beneficiary = AxisBankBeneficiary.objects.filter(
#                 account_number=beneficiary_account
#             ).first()
#             if not existing_beneficiary:
#                 AxisBankBeneficiary.objects.create(
#                     name=data.get("beneficiary_name"),
#                     account_number=beneficiary_account,
#                     bank_name="Union Bank",
#                     bank_address=data.get("beneficiary_branch_address"),
#                     ifsc_code=data.get("beneficiary_ifsc"),
#                     transaction_amount=data.get("amount"),
#                 )

#             # Generate Word document
#             template_path = os.path.join(settings.MEDIA_ROOT, 'UBForm.docx')
#             doc = DocxTemplate(template_path)
#             amount_value = data["amount"]
#             amount_words = num2words(amount_value, to="currency", lang="en_IN", currency="INR").upper()
#             context = {**data, "amount": str(amount_value), "amount_words": amount_words}
#             doc.render(context)

#             buffer = BytesIO()
#             doc.save(buffer)
#             buffer.seek(0)

#             response = HttpResponse(
#                 buffer.getvalue(),
#                 content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
#             )
#             response['Content-Disposition'] = 'attachment; filename="UnionBankForm_Filled.docx"'
#             response.set_cookie('fileDownload', 'true', max_age=10, path='/')
#             return response

#     else:
#         form = UnionBankForm(initial=remitter_data)

#     return render(request, 'form_filler/union_bank.html', {'form': form})



# # ---------------- Common Helpers ----------------
# from django.http import JsonResponse

# def remitter_lookup(request):
#     query = request.GET.get("q", "")
#     results = {"axis_bank": [], "union_bank": []}
#     if query:
#         axis_remitters = AxisBankRemitter.objects.filter(account_number__icontains=query)[:5]
#         for r in axis_remitters:
#             results["axis_bank"].append({
#                 "account_number": r.account_number,
#                 "name": r.name,
#                 "lei_code": getattr(r, "lei_code", ""),
#                 "address": r.address,
#                 "cheque_number": r.cheque_number,
#                 "mobile_number": r.mobile_number,
#                 "other_number": getattr(r, "other_number", ""),
#                 "branch": r.branch,
#                 "pan_no": getattr(r, "pan_no", ""),
#             })

#         union_remitters = UnionBankRemitter.objects.filter(account_number__icontains=query)[:5]
#         for r in union_remitters:
#             results["union_bank"].append({
#                 "account_number": r.account_number,
#                 "name": r.name,
#                 "address": r.address,
#                 "cheque_number": r.cheque_number,
#                 "mobile_number": r.mobile_number,
#                 "branch": r.branch,
#                 "pan_no": getattr(r, "pan_number", ""),
#             })

#     return JsonResponse(results)

# # ✅ Fetch beneficiary list
# def fetch_beneficiary(request):
#     q = request.GET.get('q', '')
#     if not q:
#         return JsonResponse({'results': []})

#     beneficiaries = AxisBankBeneficiary.objects.filter(name__icontains=q)[:10]
#     results = [
#         {
#             'name': b.name,
#             'account_number': b.account_number,
#             'bank_name': b.bank_name,
#             'bank_address': b.bank_address,
#             'lei_code': b.lei_code,
#             'ifsc_code': b.ifsc_code,
#         }
#         for b in beneficiaries
#     ]
#     return JsonResponse({'results': results})




from django.shortcuts import render, redirect
from django.http import HttpResponse, JsonResponse
from io import BytesIO
import os
from num2words import num2words
from docxtpl import DocxTemplate
from .forms import UnionBankForm
from .models import RemitterAccount, AxisBankBeneficiary, AxisBankRemitter, UnionBankRemitter
from django.conf import settings


def union_bank(request):
    unique_code = request.session.get('unique_code')
    remitter_data = {}

    # Prefill form if remitter exists
    if unique_code:
        remitter_acc = RemitterAccount.objects.filter(
            unique_code=unique_code,
            bank_name__icontains="Union Bank"
        ).first()

        if remitter_acc:
            remitter_data = {
                "remitter_name": remitter_acc.remitter_name or "",
                "remitter_account_number": remitter_acc.remitter_account_number or "",
                "remitter_address": remitter_acc.remitter_address or "",
                "branch": remitter_acc.branch or "",
                "pan_number": remitter_acc.pan_no or "",
                "mobile_number": remitter_acc.mobile_number or "",
                "other_number": remitter_acc.other_number or "",
            }

    if request.method == "POST":
        form = UnionBankForm(request.POST)
        if form.is_valid():
            data = form.cleaned_data

            # --- Save remitter if not already saved ---
            existing_remitter = RemitterAccount.objects.filter(
                remitter_account_number=data.get("remitter_account_number"),
                bank_name__icontains="Union Bank"
            ).first()

            if not existing_remitter:
                RemitterAccount.objects.create(
                    unique_code=unique_code,
                    bank_name="Union Bank",
                    remitter_account_number=data.get("remitter_account_number"),
                    remitter_name=data.get("remitter_name", ""),
                    remitter_address=data.get("remitter_address", ""),
                    mobile_number=data.get("mobile_number", ""),
                    pan_no=data.get("pan_number", ""),
                    branch=data.get("branch", ""),
                    other_number=data.get("other_number", "")
                )

            # --- Save beneficiary if new ---
            beneficiary_account = data.get("beneficiary_account_number")
            existing_beneficiary = AxisBankBeneficiary.objects.filter(
                account_number=beneficiary_account
            ).first()

            if not existing_beneficiary:
                AxisBankBeneficiary.objects.create(
                    unique_code=unique_code,  # ✅ FIXED: Attach unique code here
                    name=data.get("beneficiary_name"),
                    account_number=beneficiary_account,
                    bank_name="Union Bank",
                    bank_address=data.get("beneficiary_branch_address"),
                    ifsc_code=data.get("beneficiary_ifsc"),
                    # transaction_amount=data.get("amount"),
                )

                print("✅ Beneficiary saved with unique_code:", unique_code)

            # --- Generate Word document ---
            template_path = os.path.join(settings.MEDIA_ROOT, 'UBForm.docx')
            doc = DocxTemplate(template_path)
            amount_value = data["amount"]
            amount_words = num2words(amount_value, to="currency", lang="en_IN", currency="INR").upper()

            context = {**data, "amount": str(amount_value), "amount_words": amount_words}
            doc.render(context)

            buffer = BytesIO()
            doc.save(buffer)
            buffer.seek(0)

            response = HttpResponse(
                buffer.getvalue(),
                content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )
            response['Content-Disposition'] = 'attachment; filename="UnionBankForm_Filled.docx"'
            response.set_cookie('fileDownload', 'true', max_age=10, path='/')
            return response

    else:
        form = UnionBankForm(initial=remitter_data)

    return render(request, 'form_filler/union_bank.html', {'form': form})


# ---------------- Common Helpers ----------------

def remitter_lookup(request):
    query = request.GET.get("q", "")
    results = {"axis_bank": [], "union_bank": []}
    if query:
        axis_remitters = AxisBankRemitter.objects.filter(account_number__icontains=query)[:5]
        for r in axis_remitters:
            results["axis_bank"].append({
                "account_number": r.account_number,
                "name": r.name,
                "lei_code": getattr(r, "lei_code", ""),
                "address": r.address,
                "cheque_number": r.cheque_number,
                "mobile_number": r.mobile_number,
                "other_number": getattr(r, "other_number", ""),
                "branch": r.branch,
                "pan_no": getattr(r, "pan_no", ""),
            })

        union_remitters = UnionBankRemitter.objects.filter(account_number__icontains=query)[:5]
        for r in union_remitters:
            results["union_bank"].append({
                "account_number": r.account_number,
                "name": r.name,
                "address": r.address,
                "cheque_number": r.cheque_number,
                "mobile_number": r.mobile_number,
                "branch": r.branch,
                "pan_no": getattr(r, "pan_number", ""),
            })

    return JsonResponse(results)


# ✅ Fetch beneficiary list
# def fetch_beneficiary(request):
#     q = request.GET.get('q', '')
#     if not q:
#         return JsonResponse({'results': []})

#     beneficiaries = AxisBankBeneficiary.objects.filter(name__icontains=q)[:10]
#     results = [
#         {
#             'name': b.name,
#             'account_number': b.account_number,
#             'bank_name': b.bank_name,
#             'bank_address': b.bank_address,
#             # 'lei_code': b.lei_code,
#             'ifsc_code': b.ifsc_code,
#         }
#         for b in beneficiaries
#     ]
#     return JsonResponse({'results': results})








def list_docx_placeholders(request):
    template_path = os.path.join(settings.MEDIA_ROOT, 'Axis Bank Rtgs Form.docx')
    doc = DocxTemplate(template_path)
    placeholders = doc.get_undeclared_template_variables()

    html = "<h2>Available Placeholders in the Template:</h2><ul>"
    for p in sorted(placeholders):
        html += f"<li><code>{p}</code></li>"
    html += "</ul>"

    return HttpResponse(html)


def select_bank(request):
    return render(request, 'select_bank.html')

from django.shortcuts import render, redirect

# Landing page view
def landing_page(request):
    return render(request, "sign_up.html")  # your landing page template




from django.shortcuts import render, redirect
from .models import RemitterAccount
import re

def add_account(request):
    unique_code = request.session.get('unique_code')
    if not unique_code:
        print("⚠️ No session unique_code found! Redirecting to login.")
        return redirect('login')

    if request.method == 'POST':
        print("🧾 POST DATA:", request.POST)

        # Collect form data
        remitter_name = request.POST.get('remitter_name')
        remitter_account_number = request.POST.get('remitter_account_number')
        bank_name = request.POST.get('bank_name')
        branch = request.POST.get('branch')
        remitter_address = request.POST.get('remitter_address')
        mobile_number = request.POST.get('mobile_number')
        other_number = request.POST.get('other_number')
        pan_no = request.POST.get('pan_no')
        ifsc_code = request.POST.get('ifsc_code')

        errors = {}

        # ✅ Validate mobile_number
        if not re.fullmatch(r'\d{10}', mobile_number):
            errors['mobile_number'] = "Mobile Number must contain exactly 10 digits!"

        # ✅ Validate other_number
        # if not re.fullmatch(r'\d{10}', other_number):
        #     errors['other_number'] = "Other Number must contain exactly 10 digits!"

        # ✅ Validate PAN number
        if not pan_no or len(pan_no.strip()) != 10:
            errors['pan_no'] = "PAN Number must be exactly 10 characters!"

        # ✅ Validate IFSC code
        if not re.fullmatch(r'\w{11}', ifsc_code):
            errors['ifsc_code'] = "IFSC code must contain exactly 11 characters!"

        if errors:
            return render(request, 'add_account.html', {
                'errors': errors,
                'remitter_name': remitter_name,
                'remitter_account_number': remitter_account_number,
                'bank_name': bank_name,
                'branch': branch,
                'remitter_address': remitter_address,
                'mobile_number': mobile_number,
                'other_number': other_number,
                'pan_no': pan_no,
                'ifsc_code': ifsc_code
            })

        # Save to DB
        RemitterAccount.objects.create(
            unique_code=unique_code,
            remitter_name=remitter_name,
            remitter_account_number=remitter_account_number,
            bank_name=bank_name,
            branch=branch,
            remitter_address=remitter_address,
            mobile_number=mobile_number,
            other_number=other_number,
            pan_no=pan_no,
            ifsc_code=ifsc_code,
        )

        print(f"✅ Account created successfully for unique_code: {unique_code}")
        return redirect('select_bank')

    # GET request — prefill bank name if available
    bank_name = request.GET.get('bank_name', '')
    return render(request, 'add_account.html', {'bank_name': bank_name})




from django.shortcuts import render, redirect
from django.contrib.auth.hashers import make_password, check_password
from .models import SignUp
import uuid
import re

# ------------------ SIGN UP ------------------ #
from django.shortcuts import render, redirect
from .models import SignUp
import re
import uuid

def sign_up(request):
    if request.method == 'POST':
        username = request.POST.get('username')
        email = request.POST.get('email')
        phone = request.POST.get('phone')
        password = request.POST.get('password')
        confirm_password = request.POST.get('confirm_password')

        errors = {}

        if password != confirm_password:
            errors['confirm_password'] = "Passwords do not match!"
        if len(password) < 8:
            errors['password'] = "Password must be at least 8 characters long!"
        if not re.fullmatch(r'\d{10}', phone):
            errors['phone'] = "Phone number must contain exactly 10 digits!"
        if SignUp.objects.filter(email=email).exists():
            errors['email'] = "Email already registered!"

        if errors:
            return render(request, 'sign_up.html', {
                'errors': errors,
                'username': username,
                'email': email,
                'phone': phone
            })

        # Create user (password stored as plain text)
        user = SignUp.objects.create(
            username=username,
            email=email,
            phone=phone,
            password=password  # ⚠️ plain text for testing only
        )

        request.session['unique_code'] = user.unique_code
        return redirect('select_bank')

    return render(request, 'sign_up.html')



#from django.shortcuts import render, redirect
from django.shortcuts import render, redirect
from .models import SignUp

def login_view(request):
    if request.method == 'POST':
        email = request.POST.get('email', '').strip()
        password = request.POST.get('password', '').strip()

        try:
            user = SignUp.objects.get(email=email)
        except SignUp.DoesNotExist:
            return render(request, 'login.html', {
                'errors': {'email': "Invalid email or password"},
                'email': email
            })

        if password == user.password:
            request.session['unique_code'] = user.unique_code
            return redirect('select_bank')
        else:
            return render(request, 'login.html', {
                'errors': {'password': "Invalid email or password"},
                'email': email
            })

    return render(request, 'login.html')







from django.shortcuts import redirect
from django.contrib import messages

def logout_view(request):
    """Logs out the user and clears the session."""
    request.session.flush()  # clears all session data
    messages.success(request, "You have been logged out successfully.")
    print("👋 User logged out and session cleared.")
    return redirect('login')




from django.shortcuts import render, redirect
from .models import RemitterAccount

# def select_bank(request):
#     # ✅ Ensure user is logged in
#     unique_code = request.session.get('unique_code')

#     if not unique_code:
#         print("⚠️ No session found — redirecting to login")
#         return redirect('login')

#     # ✅ Fetch all accounts linked to the logged-in user's unique_code
#     accounts = RemitterAccount.objects.filter(unique_code=unique_code)
#     print(f"✅ Showing {accounts.count()} accounts for unique_code: {unique_code}")

#     # ✅ Render the select_bank page with those accounts
#     return render(request, 'select_bank.html', {'accounts': accounts})



def select_bank(request):
    # ✅ Try to get from session first
    unique_code = request.session.get('unique_code')

    # ✅ If missing, try to get from user or DB and store in session
    if not unique_code:
        # Example: You might be saving unique_code at login or during signup
        # Adjust this logic depending on your app
        if hasattr(request.user, 'unique_code'):
            unique_code = request.user.unique_code
        else:
            # Fallback: generate a temporary code or handle error
            print("⚠️ No unique_code found for user — generating temporary code")
            import uuid
            unique_code = str(uuid.uuid4())[:8]  # short temp ID

        request.session['unique_code'] = unique_code
        print(f"✅ Stored unique_code in session: {unique_code}")

    # ✅ Now guaranteed to have it
    accounts = RemitterAccount.objects.filter(unique_code=unique_code)
    print(f"✅ Showing {accounts.count()} accounts for unique_code: {unique_code}")

    return render(request, 'select_bank.html', {'accounts': accounts})





from django.shortcuts import render, redirect
from django.contrib import messages
from .models import RemitterAccount

def save_bank_account(request, bank_name):
    if request.method == 'POST':
        print("🧾 POST DATA:", request.POST)  # ✅ Debug: see what’s really coming

        # Correct field names
        remitter_account_number = (
            request.POST.get('remitter_account_number') or
            request.POST.get('account_number')  # fallback if name mismatch
        )
        remitter_name = request.POST.get('remitter_name') or request.POST.get('name')
        remitter_address = request.POST.get('remitter_address') or request.POST.get('address')

        if not remitter_account_number:
            messages.error(request, "Remitter Account Number is missing!")
            return redirect('select_bank_2')

        RemitterAccount.objects.create(
            unique_code=request.session.get('unique_code'),
            bank_name=bank_name,
            remitter_account_number=remitter_account_number,
            remitter_name=remitter_name,
            branch=request.POST.get('branch'),
            pan_no=request.POST.get('pan_no'),
            remitter_address=remitter_address,
            mobile_number=request.POST.get('mobile_number'),
            other_number=request.POST.get('other_number'),
        )

        messages.success(request, f"{bank_name} account added successfully!")
        return redirect('select_bank_2')  # ✅ redirect, not render






from django.shortcuts import render, redirect
from .models import RemitterAccount

def bank_accounts(request):
    unique_code = request.session.get("unique_code")
    print("🧭 Session unique_code:", unique_code)
    if not unique_code:
        return redirect("login")

    accounts = RemitterAccount.objects.filter(unique_code=unique_code)
    bank_name = accounts.first().bank_name if accounts.exists() else "Your"

    return render(request, "bank_accounts.html", {
        "accounts": accounts,
        "bank_name": bank_name,
    })




from django.shortcuts import render, redirect
from .models import RemitterAccount

def select_bank_2(request):
    # ✅ Ensure user is logged in
    unique_code = request.session.get('unique_code')

    if not unique_code:
        print("⚠️ No session found — redirecting to login")
        return redirect('login')

    # ✅ Fetch all accounts for this user
    accounts = RemitterAccount.objects.filter(unique_code=unique_code)
    print(f"✅ Showing {accounts.count()} accounts for unique_code: {unique_code}")

    # ✅ Render the select_bank_2.html page
    return render(request, 'select_bank_2.html', {'accounts': accounts})




from django.http import HttpResponse
import openpyxl

def download_beneficiaries_excel(request):
    """Download all beneficiaries for the logged-in user as an Excel file."""
    unique_code = request.session.get('unique_code')
    if not unique_code:
        return redirect('login')

    beneficiaries = AxisBankBeneficiary.objects.filter(unique_code=unique_code)

    # Create Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Beneficiaries"

    # Header row
    headers = [
        "Name", "Account Number", "Bank Name", "Bank Address",
        "IFSC Code"
    ]
    ws.append(headers)

    # Data rows
    for b in beneficiaries:
        ws.append([
            b.name, b.account_number, b.bank_name,
            b.bank_address or "", b.ifsc_code or "",
        ])

    # Prepare HTTP response
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="beneficiaries.xlsx"'
    wb.save(response)
    return response




from django.shortcuts import render, redirect
from .models import AxisBankBeneficiary

def beneficiaries_list(request):
    """Show all beneficiaries for the logged-in user."""
    unique_code = request.session.get('unique_code')
    if not unique_code:
        return redirect('login')

    beneficiaries = AxisBankBeneficiary.objects.filter(unique_code=unique_code)
    return render(request, 'beneficiaries_list.html', {'beneficiaries': beneficiaries})





# views.py
def select_specific_bank(request, bank_name):
    unique_code = request.session.get('unique_code')

    if not unique_code:
        print("⚠️ No session found — redirecting to login")
        return redirect('login')

    # Normalize and filter
    bank_name = bank_name.strip().lower()
    accounts = RemitterAccount.objects.filter(
        unique_code=unique_code,
        bank_name__icontains=bank_name
    )

    print(f"✅ Showing {accounts.count()} {bank_name.title()} accounts for unique_code: {unique_code}")

    return render(request, 'bank_specific.html', {
        'accounts': accounts,
        'bank_name': bank_name.title()
    })






from django.shortcuts import render, get_object_or_404, redirect
from .models import AxisBankBeneficiary  # or your actual model
from django.contrib import messages

def beneficiaries_list(request):
    unique_code = request.session.get('unique_code')
    beneficiaries = AxisBankBeneficiary.objects.filter(unique_code=unique_code)
    return render(request, 'beneficiaries_list.html', {'beneficiaries': beneficiaries})


def edit_beneficiary(request, beneficiary_id):
    beneficiary = get_object_or_404(AxisBankBeneficiary, id=beneficiary_id)

    if request.method == 'POST':
        beneficiary.name = request.POST.get('name')
        beneficiary.account_number = request.POST.get('account_number')
        beneficiary.bank_name = request.POST.get('bank_name')
        beneficiary.bank_address = request.POST.get('bank_address')
        beneficiary.ifsc_code = request.POST.get('ifsc_code')
        beneficiary.save()
        messages.success(request, "Beneficiary updated successfully!")
        return redirect('beneficiaries_list')

    return render(request, 'edit_beneficiary.html', {'beneficiary': beneficiary})


def delete_beneficiary(request, beneficiary_id):
    beneficiary = get_object_or_404(AxisBankBeneficiary, id=beneficiary_id)
    beneficiary.delete()
    messages.success(request, "Beneficiary deleted successfully!")
    return redirect('beneficiaries_list')





from django.shortcuts import render

def select_bank_3(request):
    return render(request, 'select_bank_3.html')




def hdfc_bank_view(request):
    return render(request, 'form_filler/hdfc_bank.html')
